from io import BytesIO
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analysis.top2 import concept_metric_columns, top2_min_for
from .parsing import dataframe_columns


def _data_letter(df: pd.DataFrame, structure: dict, name: str):
    headers = dataframe_columns(structure)
    try:
        actual = headers[headers.index(name)]
    except ValueError:
        return None
    try:
        return get_column_letter(list(df.columns).index(actual) + 1)
    except ValueError:
        return None


def _frame_letter(df: pd.DataFrame, name: str):
    try:
        return get_column_letter(list(df.columns).index(name) + 1)
    except ValueError:
        return None


def _add_data_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    sheet = wb.create_sheet(title)
    sheet.append(list(df.columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in df.itertuples(index=False):
        sheet.append([_safe_value(value) for value in row])
    sheet.freeze_panes = "A2"

FLAG_FILLS = {
    "выше": PatternFill("solid", fgColor="D8F0E2"),
    "ниже": PatternFill("solid", fgColor="FDEAEA"),
    "не значимо": PatternFill("solid", fgColor="F2F4F6"),
    "нет данных": PatternFill("solid", fgColor="FFFFFF"),
}


def _safe_value(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, bool):
        return bool(value)
    if isinstance(value, (int, float)):
        return value
    return str(value)


def _add_frequencies_sheet(wb: Workbook, df: pd.DataFrame, column: str, structure: dict) -> None:
    sheet = wb.create_sheet("Частоты")
    letter = get_column_letter(list(df.columns).index(column) + 1)
    data_range = f"'Данные'!${letter}$2:${letter}${len(df) + 1}"
    column_meta = next((c for c in structure["columns"] if c["name"] == column), None)
    title = column
    if column_meta and column_meta.get("label"):
        title = f"{column} — {column_meta['label']}"
    sheet.append([title])
    sheet["A1"].font = Font(bold=True)
    sheet.append([])
    sheet.append(["Значение", "N", "% от валидных"])
    for cell in sheet[3]:
        cell.font = Font(bold=True)

    counts = df[column].value_counts(dropna=True)
    start_row = 4
    for i, value in enumerate(counts.index):
        row = start_row + i
        sheet.append([_safe_value(value)])
        sheet.cell(row=row, column=2, value=f"=COUNTIF({data_range},A{row})")
        total_row = start_row + len(counts)
        sheet.cell(row=row, column=3, value=f"=IF($B${total_row}=0,0,B{row}/$B${total_row})")
        sheet.cell(row=row, column=3).number_format = "0.0%"
    total_row = start_row + len(counts)
    sheet.cell(row=total_row, column=1, value="Итого валидных").font = Font(bold=True)
    sheet.cell(row=total_row, column=2, value=f"=SUM(B{start_row}:B{total_row - 1})").font = Font(bold=True)

    numeric = pd.to_numeric(df[column], errors="coerce")
    if int(numeric.notna().sum()) >= 2:
        r = total_row + 2
        sheet.cell(row=r, column=1, value="Числовые показатели").font = Font(bold=True)
        metrics = [
            ("Среднее", f"=AVERAGE({data_range})"),
            ("Медиана", f"=MEDIAN({data_range})"),
            ("Стандартное отклонение", f"=STDEV.S({data_range})"),
            ("Минимум", f"=MIN({data_range})"),
            ("Максимум", f"=MAX({data_range})"),
            ("N (числовых)", f"=COUNT({data_range})"),
        ]
        for j, (label, formula) in enumerate(metrics, start=1):
            sheet.cell(row=r + j, column=1, value=label)
            sheet.cell(row=r + j, column=2, value=formula)


def _add_top2_sheet(wb: Workbook, df: pd.DataFrame, structure: dict, config: dict, result: dict) -> None:
    norm_mode = result.get("norm_mode", "q3")
    norm_labels = {
        "q3": "Q3 по концептам файла",
        "median": "медиана по концептам файла",
        "manual": "фиксированная норма (задана вручную)",
        "pool_q3": "Q3 по пулу волн (текущая + предыдущая)",
        "pool_median": "медиана по пулу волн (текущая + предыдущая)",
    }
    sheet = wb.create_sheet("Top2 и нормы")
    sheet["A1"] = "Top2% (верхние 2 категории шкалы) и 95% интервал Уилсона"
    sheet["A1"].font = Font(bold=True)
    sheet["A2"] = f"Норма: {norm_labels.get(norm_mode, norm_mode)}"
    sheet["A2"].font = Font(italic=True)
    sheet["K1"] = "z (95%)"
    sheet["L1"] = 1.959964

    headers = ["Концепт", "Метрика", "N", "Top2", "%", "Нижняя 95%", "Верхняя 95%", "Норма", "Вывод"]
    header_row = 3
    for i, title in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=i, value=title)
        cell.font = Font(bold=True)

    triples = concept_metric_columns(structure, result["metrics"], config.get("concept_map", {}))
    column_letters = {name: get_column_letter(list(df.columns).index(name) + 1) for _, _, name in triples}
    last_data_row = len(df) + 1

    row = header_row + 1
    for metric in result["metrics"]:
        metric_rows = [(concept, col) for concept, m, col in triples if m == metric]
        metric_start = row
        metric_end = row + len(metric_rows) - 1
        for concept_name, column_name in metric_rows:
            letter = column_letters[column_name]
            data_range = f"'Данные'!${letter}$2:${letter}${last_data_row}"
            minimum = top2_min_for(structure, column_name)
            sheet.cell(row=row, column=1, value=concept_name)
            sheet.cell(row=row, column=2, value=metric)
            sheet.cell(row=row, column=3, value=f"=COUNT({data_range})")
            sheet.cell(row=row, column=4, value=f'=COUNTIF({data_range},">={minimum}")')
            sheet.cell(row=row, column=5, value=f'=IF(C{row}=0,"",D{row}/C{row})')
            wilson_base = f"(E{row}+$L$1^2/(2*C{row}))/(1+$L$1^2/C{row})"
            wilson_margin = f"$L$1*SQRT(E{row}*(1-E{row})/C{row}+$L$1^2/(4*C{row}^2))/(1+$L$1^2/C{row})"
            sheet.cell(row=row, column=6, value=f'=IF(OR(C{row}=0,E{row}=""),"",{wilson_base}-{wilson_margin})')
            sheet.cell(row=row, column=7, value=f'=IF(OR(C{row}=0,E{row}=""),"",{wilson_base}+{wilson_margin})')
            if norm_mode == "q3":
                sheet.cell(
                    row=row,
                    column=8,
                    value=f'=IF(COUNT(F{metric_start}:F{metric_end})=0,"",QUARTILE.INC(F{metric_start}:F{metric_end},3))',
                )
            elif norm_mode == "median":
                sheet.cell(
                    row=row,
                    column=8,
                    value=f'=IF(COUNT(F{metric_start}:F{metric_end})=0,"",MEDIAN(F{metric_start}:F{metric_end}))',
                )
            else:
                norm_value = result["norms"].get(metric)
                sheet.cell(row=row, column=8, value=None if norm_value is None else norm_value / 100)
            sheet.cell(
                row=row,
                column=9,
                value=f'=IF(OR(F{row}="",H{row}=""),"",IF(F{row}>H{row},"выше",IF(G{row}<H{row},"ниже","не значимо")))',
            )
            for col in (5, 6, 7, 8):
                sheet.cell(row=row, column=col).number_format = "0.0%"
            flag = result["cells"].get(concept_name, {}).get(metric, {}).get("flag")
            sheet.cell(row=row, column=9).fill = FLAG_FILLS.get(flag, FLAG_FILLS["нет данных"])
            row += 1


def _add_wave_sheet(wb: Workbook, task: dict, df: pd.DataFrame, structure: dict) -> None:
    result = task["result"]
    frames = task.get("frames", {})
    mode = result["mode"]
    alpha = float(result.get("alpha", 0.05))
    mode_labels = {
        "independent": "полные выборки (независимые)",
        "panel": "панель — одни и те же респонденты",
        "panel_lag": "панель с окном по времени между волнами",
    }
    sheet = wb.create_sheet("Сравнение волн")
    sheet["A1"] = "Сравнение волн: Top2% (верхние 2 категории шкалы), разница и значимость"
    sheet["A1"].font = Font(bold=True)
    sheet["A2"] = f"Режим: {mode_labels.get(mode, mode)} · α = {alpha}"
    sheet["A2"].font = Font(italic=True)
    if result.get("matched_total") is not None:
        sheet["A3"] = (
            f"Совпало ID: {result['matched_total']} · в расчёте: {result['matched_used']}"
            + (f" · исключено окном лага: {result['lag']['excluded']}" if result.get("lag") else "")
        )
        sheet["A3"].font = Font(italic=True)
    sheet["N1"] = "z крит."
    sheet["O1"] = f"=NORM.S.INV(1-{alpha}/2)"

    headers = [
        "Концепт",
        "Метрика",
        "N волна 1",
        "Top2 волна 1",
        "N волна 2",
        "Top2 волна 2",
        "Разница, п.п.",
        "Нижняя 95%",
        "Верхняя 95%",
        "p-value",
        "Вывод",
    ]
    header_row = 5
    for i, title in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=i, value=title)
        cell.font = Font(bold=True)

    row = header_row + 1
    for concept in result["concepts"]:
        concept_cells = result["cells"].get(concept, {})
        if not concept_cells:
            continue
        first_row = row
        for metric in result["metrics"]:
            cell = concept_cells.get(metric)
            if not cell:
                continue
            if mode == "independent":
                prev_df = frames.get("prev")
                main_letter = _data_letter(df, structure, cell["main_column"])
                prev_letter = _frame_letter(prev_df, cell["prev_column"]) if prev_df is not None else None
                main_min = cell.get("main_min", 4)
                prev_min = cell.get("prev_min", 4)
                main_range = f"'Данные'!${main_letter}$2:${main_letter}${len(df) + 1}" if main_letter else None
                prev_range = f"'Данные (волна 2)'!${prev_letter}$2:${prev_letter}${len(prev_df) + 1}" if prev_letter and prev_df is not None else None
                if main_range and prev_range:
                    sheet.cell(row=row, column=3, value=f"=COUNT({main_range})")
                    sheet.cell(row=row, column=4, value=f'=IF(C{row}=0,"",COUNTIF({main_range},">={main_min}")/C{row})')
                    sheet.cell(row=row, column=5, value=f"=COUNT({prev_range})")
                    sheet.cell(row=row, column=6, value=f'=IF(E{row}=0,"",COUNTIF({prev_range},">={prev_min}")/E{row})')
                    sheet.cell(row=row, column=7, value=f'=IF(OR(D{row}="",F{row}=""),"",D{row}-F{row})')
                    se = f"SQRT(D{row}*(1-D{row})/C{row}+F{row}*(1-F{row})/E{row})"
                    sheet.cell(row=row, column=8, value=f'=IF(G{row}="","",G{row}-$O$1*{se})')
                    sheet.cell(row=row, column=9, value=f'=IF(G{row}="","",G{row}+$O$1*{se})')
                    sheet.cell(
                        row=row,
                        column=10,
                        value=f'=IF(G{row}="","",2*(1-NORM.S.DIST(ABS(G{row}/{se}),TRUE)))',
                    )
                    sheet.cell(
                        row=row,
                        column=11,
                        value=f'=IF(OR(H{row}="",I{row}=""),"",IF(H{row}>0,"выше",IF(I{row}<0,"ниже","не значимо")))',
                    )
            else:
                merged = frames.get("merged")
                main_letter = _frame_letter(merged, cell["main_column"]) if merged is not None else None
                prev_letter = _frame_letter(merged, cell["prev_column"]) if merged is not None else None
                if main_letter and prev_letter and merged is not None:
                    last = len(merged) + 1
                    main_range = f"'Пары волн'!${main_letter}$2:${main_letter}${last}"
                    prev_range = f"'Пары волн'!${prev_letter}$2:${prev_letter}${last}"
                    main_min = cell.get("main_min", 4)
                    prev_min = cell.get("prev_min", 4)
                    b_expr = f'COUNTIFS({main_range},">={main_min}",{prev_range},"<{prev_min}")'
                    c_expr = f'COUNTIFS({main_range},"<{main_min}",{prev_range},">={prev_min}")'
                    sheet.cell(row=row, column=3, value=f"=COUNT({main_range})")
                    sheet.cell(row=row, column=4, value=f'=IF(C{row}=0,"",COUNTIF({main_range},">={main_min}")/C{row})')
                    sheet.cell(row=row, column=5, value=f"=COUNT({prev_range})")
                    sheet.cell(row=row, column=6, value=f'=IF(E{row}=0,"",COUNTIF({prev_range},">={prev_min}")/E{row})')
                    sheet.cell(row=row, column=7, value=f'=IF(OR(D{row}="",F{row}=""),"",D{row}-F{row})')
                    se = f"SQRT(MAX({b_expr}+{c_expr}-({b_expr}-{c_expr})^2/C{row},0))/C{row}"
                    sheet.cell(row=row, column=8, value=f'=IF(G{row}="","",G{row}-$O$1*{se})')
                    sheet.cell(row=row, column=9, value=f'=IF(G{row}="","",G{row}+$O$1*{se})')
                    sheet.cell(
                        row=row,
                        column=10,
                        value=f'=IF({b_expr}+{c_expr}=0,1,CHISQ.DIST.RT((ABS({b_expr}-{c_expr})-1)^2/({b_expr}+{c_expr}),1))',
                    )
                    sheet.cell(
                        row=row,
                        column=11,
                        value=f'=IF(OR(H{row}="",I{row}=""),"",IF(H{row}>0,"выше",IF(I{row}<0,"ниже","не значимо")))',
                    )
            sheet.cell(row=row, column=1, value=concept)
            sheet.cell(row=row, column=2, value=metric)
            for col in (4, 6, 8, 9):
                sheet.cell(row=row, column=col).number_format = "0.0%"
            sheet.cell(row=row, column=7).number_format = "+0.0%;-0.0%"
            sheet.cell(row=row, column=10).number_format = "0.0000"
            flag = cell.get("flag")
            sheet.cell(row=row, column=11).fill = FLAG_FILLS.get(flag, FLAG_FILLS["нет данных"])
            row += 1
        if row - first_row > 1:
            sheet.merge_cells(start_row=first_row, start_column=1, end_row=row - 1, end_column=1)
            sheet.cell(row=first_row, column=1).alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A6"


def build_project_export(
    project_name: str,
    df: pd.DataFrame,
    structure: dict,
    history: list,
    column: Optional[str] = None,
    task: Optional[dict] = None,
) -> bytes:
    wb = Workbook()
    data_sheet = wb.active
    data_sheet.title = "Данные"
    data_sheet.append(list(df.columns))
    for cell in data_sheet[1]:
        cell.font = Font(bold=True)
    for row in df.itertuples(index=False):
        data_sheet.append([_safe_value(value) for value in row])
    data_sheet.freeze_panes = "A2"

    steps_sheet = wb.create_sheet("Шаги")
    steps_sheet.append(["№", "Действие", "Строк до", "Строк после", "Когда"])
    for cell in steps_sheet[1]:
        cell.font = Font(bold=True)
    if history:
        for i, step in enumerate(history, start=1):
            steps_sheet.append([i, step["summary"], step["rows_before"], step["rows_after"], step["created_at"]])
    else:
        steps_sheet.append([0, "Исходные данные (без шагов обработки)", len(df), len(df), ""])
    steps_sheet.append([])
    steps_sheet.append(["Проект", project_name])

    if column and column in df.columns:
        _add_frequencies_sheet(wb, df, column, structure)

    if task and task.get("task") == "top2_norms":
        _add_top2_sheet(wb, df, structure, task.get("config", {}), task.get("result", {}))
    if task and task.get("task") == "wave_compare":
        frames = task.get("frames", {})
        if frames.get("prev") is not None:
            _add_data_sheet(wb, "Данные (волна 2)", frames["prev"])
        if frames.get("merged") is not None:
            _add_data_sheet(wb, "Пары волн", frames["merged"])
        _add_wave_sheet(wb, task, df, structure)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
